from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable, Iterator

import nibabel as nib
import numpy as np
from openpyxl import load_workbook

from core.exceptions import InvalidInputError
from data.structures import CompetitionDataset, Series, Study


_NIFTI_SUFFIXES = (".nii", ".nii.gz")
_MASK_HINTS = ("mask", "seg", "label", "roi")
_SERIES_TYPE_HEADERS = ("accessionnumber", "seriesuid", "seriestype")


def _nifti_stem(path: Path) -> str:
    return path.name[:-7] if path.name.lower().endswith(".nii.gz") else path.stem


def _select_original_nifti_files(root: Path, files: Iterable[Path]) -> list[Path]:
    selected: list[Path] = []
    grouped: dict[Path, list[Path]] = defaultdict(list)
    for path in files:
        if len(path.relative_to(root).parts) <= 2:
            selected.append(path)
        else:
            grouped[path.parent].append(path)

    for directory, paths in grouped.items():
        if len(paths) == 1:
            selected.extend(paths)
            continue
        originals = [path for path in paths if _nifti_stem(path) == directory.name]
        if len(originals) != 1:
            raise InvalidInputError(
                f"series directory {directory} has multiple NIfTI files but "
                f"expected exactly one original named {directory.name}.nii or "
                f"{directory.name}.nii.gz: {[path.name for path in sorted(paths)]}"
            )
        selected.extend(originals)
    return sorted(selected)


def _clean_identifier(value: Any, fallback: str) -> str:
    text = str(value or "").strip()
    if not text:
        text = fallback
    return re.sub(r"[\\/\x00-\x1f]", "_", text)


def _metadata_key(value: Any) -> str:
    return re.sub(r"\s+", "", str(value if value is not None else "")).casefold()


def _read_series_types(root: Path) -> dict[tuple[str, str], str]:
    path = root / "SeriesType.xlsx"
    if not path.is_file():
        return {}

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise InvalidInputError(
            f"cannot read series metadata {path}: {type(exc).__name__}: {exc}"
        ) from exc

    try:
        header: tuple[Any, int, dict[str, int]] | None = None
        for worksheet in workbook.worksheets:
            for row_number, row in enumerate(
                worksheet.iter_rows(values_only=True),
                start=1,
            ):
                cells = [_metadata_key(value) for value in row]
                if all(name in cells for name in _SERIES_TYPE_HEADERS):
                    header = (
                        worksheet,
                        row_number,
                        {name: cells.index(name) for name in _SERIES_TYPE_HEADERS},
                    )
                    break
            if header is not None:
                break

        if header is None:
            raise InvalidInputError(
                f"series metadata {path} is missing headers: "
                "AccessionNumber, SeriesUid, SeriesType"
            )

        worksheet, header_row, columns = header
        series_types: dict[tuple[str, str], str] = {}
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            values = {
                name: row[index] if index < len(row) else None
                for name, index in columns.items()
            }
            if all(
                not str(value if value is not None else "").strip()
                for value in values.values()
            ):
                continue
            missing = [
                name
                for name, value in values.items()
                if not str(value if value is not None else "").strip()
            ]
            if missing:
                raise InvalidInputError(
                    f"series metadata {path} sheet={worksheet.title!r} "
                    f"row={row_number} is missing {', '.join(missing)}"
                )

            key = (
                _metadata_key(values["accessionnumber"]),
                _metadata_key(values["seriesuid"]),
            )
            series_type = str(values["seriestype"]).strip()
            previous = series_types.get(key)
            if previous is not None and previous != series_type:
                raise InvalidInputError(
                    f"series metadata {path} has conflicting SeriesType values "
                    f"for accession={values['accessionnumber']!r} "
                    f"series={values['seriesuid']!r}: {previous!r}, {series_type!r}"
                )
            series_types[key] = series_type
        return series_types
    except InvalidInputError:
        raise
    except Exception as exc:
        raise InvalidInputError(
            f"cannot read series metadata {path}: {type(exc).__name__}: {exc}"
        ) from exc
    finally:
        workbook.close()


class DatasetLoader:
    """Load a NIfTI tree into the competition domain model."""

    def load(self, dataset_path: str | Path) -> CompetitionDataset:
        """Compatibility entry point for callers that need the full dataset."""
        root = Path(dataset_path).expanduser().resolve()
        return CompetitionDataset(
            root,
            tuple(self.iter_studies(root)),
            {"format": "nifti"},
        )

    def iter_studies(self, dataset_path: str | Path) -> Iterator[Study]:
        """Discover file paths, then load one study at a time."""
        root = Path(dataset_path).expanduser().resolve()
        if not root.is_dir():
            raise InvalidInputError(f"dataset_path is not a directory: {root}")

        nifti_files = _select_original_nifti_files(
            root,
            (
                path
                for path in root.rglob("*")
                if path.is_file()
                and path.name.lower().endswith(_NIFTI_SUFFIXES)
                and not any(hint in path.name.lower() for hint in _MASK_HINTS)
            ),
        )
        if not nifti_files:
            raise InvalidInputError(f"no readable NIfTI images under {root}")
        yield from self._iter_nifti(root, nifti_files, _read_series_types(root))

    def _iter_nifti(
        self,
        root: Path,
        files: Iterable[Path],
        series_types: dict[tuple[str, str], str],
    ) -> Iterator[Study]:
        grouped: dict[str, list[Path]] = defaultdict(list)
        for path in files:
            relative = path.relative_to(root)
            accession = relative.parts[0] if len(relative.parts) > 1 else _nifti_stem(path)
            grouped[_clean_identifier(accession, "study")].append(path)

        for accession, paths in sorted(grouped.items()):
            yield Study(
                accession_number=accession,
                series=tuple(
                    self._read_nifti_series(
                        root,
                        accession,
                        path,
                        series_types,
                    )
                    for path in paths
                ),
            )

    def _read_nifti_series(
        self,
        root: Path,
        accession: str,
        path: Path,
        series_types: dict[tuple[str, str], str],
    ) -> Series:
        relative = path.relative_to(root)
        if len(relative.parts) == 1 or path.parent == root / relative.parts[0]:
            series_uid = _nifti_stem(path)
        else:
            series_uid = path.parent.name
        sidecar = path.with_name(_nifti_stem(path) + ".json")
        try:
            metadata: dict[str, Any] = {}
            if sidecar.is_file():
                metadata = json.loads(sidecar.read_text(encoding="utf-8"))

            image = nib.load(str(path))
            array = np.asanyarray(image.dataobj)
            array = np.squeeze(array)
            if array.ndim != 3:
                raise InvalidInputError(f"NIfTI image must be 3-D: {path} -> {array.shape}")

            uid = _clean_identifier(
                metadata.get("SeriesInstanceUID"),
                series_uid,
            )
            description = str(
                series_types.get((_metadata_key(accession), _metadata_key(uid)))
                or metadata.get("SeriesDescription")
                or metadata.get("ProtocolName")
                or series_uid
            )
            return Series(
                series_uid=uid,
                modality=description,
                image=np.asarray(array),
                affine=np.asarray(image.affine, dtype=np.float64),
                source_path=path,
                metadata=metadata,
            )
        except Exception as exc:
            detail = (
                str(exc)
                if isinstance(exc, InvalidInputError)
                else f"{type(exc).__name__}: {exc}"
            )
            raise InvalidInputError(
                f"cannot load NIfTI study={accession!r} series={series_uid!r} "
                f"path={path}: {detail}"
            ) from exc
